"""Document editing — replace, delete, insert_after with backup."""

import csv
import io
import os
import shutil
from pathlib import Path

from docwizard.config import SUPPORTED_EDIT_FORMATS


class UnsupportedEditError(Exception):
    pass


def apply_edits(path: str, edits: list[dict]) -> str:
    """Apply edits to a document.

    Supported edit types:
      replace:      {"type": "replace", "search": "...", "replace": "..."}
      delete:       {"type": "delete", "search": "..."}
      insert_after: {"type": "insert_after", "search": "...", "replace": "..."}

    Returns a summary of changes made.
    """
    p = Path(path)
    ext = p.suffix.lower()

    if ext not in SUPPORTED_EDIT_FORMATS:
        raise UnsupportedEditError(
            f"Cannot edit '{ext}' files. Editable formats: "
            f"{', '.join(sorted(SUPPORTED_EDIT_FORMATS))}"
        )

    # Create backup
    backup_path = str(path) + ".bak"
    try:
        shutil.copy2(path, backup_path)
    except OSError:
        pass  # Non-fatal

    results = []
    for edit in edits:
        edit_type = edit.get("type", "replace")
        search = edit.get("search", "")
        replace = edit.get("replace", "")

        if not search:
            results.append(f"Skipped edit with empty search string.")
            continue

        if edit_type == "delete":
            count = _dispatch_replace(ext, path, search, "")
            if count > 0:
                results.append(
                    f"Deleted '{_truncate(search, 40)}' "
                    f"({count} occurrence{'s' if count != 1 else ''})"
                )
            else:
                results.append(f"No match found for '{_truncate(search, 40)}'")

        elif edit_type == "replace":
            count = _dispatch_replace(ext, path, search, replace)
            if count > 0:
                results.append(
                    f"Replaced '{_truncate(search, 40)}' → "
                    f"'{_truncate(replace, 40)}' "
                    f"({count} occurrence{'s' if count != 1 else ''})"
                )
            else:
                results.append(f"No match found for '{_truncate(search, 40)}'")

        elif edit_type == "insert_after":
            count = _dispatch_insert_after(ext, path, search, replace)
            if count > 0:
                results.append(
                    f"Inserted content after '{_truncate(search, 40)}' "
                    f"({count} location{'s' if count != 1 else ''})"
                )
            else:
                results.append(f"No match found for '{_truncate(search, 40)}'")

        else:
            results.append(f"Skipped unknown edit type: '{edit_type}'")

    summary = "\n".join(results)
    return f"Backup saved: {backup_path}\n{summary}"


# ── Dispatch helpers ──────────────────────────────────────────────────

def _dispatch_replace(ext: str, path: str, search: str, replace: str) -> int:
    if ext == ".docx":
        return _docx_replace(path, search, replace)
    elif ext == ".csv":
        return _csv_replace(path, search, replace)
    elif ext == ".xlsx":
        return _xlsx_replace(path, search, replace)
    else:
        return _txt_replace(path, search, replace)


def _dispatch_insert_after(ext: str, path: str, search: str, content: str) -> int:
    if ext == ".docx":
        return _docx_insert_after(path, search, content)
    else:
        return _txt_insert_after(path, search, content)


# ── Plain text ────────────────────────────────────────────────────────

def _txt_replace(path: str, search: str, replace: str) -> int:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    count = content.count(search)
    if count > 0:
        content = content.replace(search, replace)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)

    return count


def _txt_insert_after(path: str, search: str, content: str) -> int:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        lines = f.readlines()

    count = 0
    new_lines = []
    for line in lines:
        new_lines.append(line)
        if search in line:
            new_lines.append(content + "\n")
            count += 1

    if count > 0:
        with open(path, "w", encoding="utf-8") as f:
            f.writelines(new_lines)

    return count


# ── CSV ───────────────────────────────────────────────────────────────

def _csv_replace(path: str, search: str, replace: str) -> int:
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        reader = csv.reader(f)
        rows = list(reader)

    count = 0
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if search in cell:
                count += cell.count(search)
                rows[i][j] = cell.replace(search, replace)

    if count > 0:
        with open(path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerows(rows)

    return count


# ── DOCX ──────────────────────────────────────────────────────────────

def _merge_runs(para) -> None:
    """Merge all runs in a paragraph into the first run, preserving its formatting.

    This ensures search strings that span multiple runs are handled correctly.
    """
    if len(para.runs) <= 1:
        return
    full_text = "".join(run.text for run in para.runs)
    para.runs[0].text = full_text
    for run in para.runs[1:]:
        run.text = ""


def _docx_replace(path: str, search: str, replace: str) -> int:
    from docx import Document

    doc = Document(path)
    count = 0

    for para in doc.paragraphs:
        if search in para.text:
            count += para.text.count(search)
            _merge_runs(para)
            para.runs[0].text = para.runs[0].text.replace(search, replace)

    if count > 0:
        doc.save(path)

    return count


def _docx_insert_after(path: str, search: str, content: str) -> int:
    from docx import Document
    from docx.oxml import OxmlElement

    doc = Document(path)
    matches = [para for para in doc.paragraphs if search in para.text]

    for para in matches:
        new_para = OxmlElement("w:p")
        new_run = OxmlElement("w:r")
        new_text = OxmlElement("w:t")
        new_text.text = content
        new_run.append(new_text)
        new_para.append(new_run)
        para._element.addnext(new_para)

    if matches:
        doc.save(path)

    return len(matches)


# ── XLSX ──────────────────────────────────────────────────────────────

def _xlsx_replace(path: str, search: str, replace: str) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    count = 0

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and search in cell.value:
                    count += cell.value.count(search)
                    cell.value = cell.value.replace(search, replace)

    if count > 0:
        wb.save(path)

    wb.close()
    return count


# ── Utility ───────────────────────────────────────────────────────────

def _truncate(text: str, max_len: int) -> str:
    if len(text) <= max_len:
        return text
    return text[: max_len - 3] + "..."
