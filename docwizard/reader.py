"""Multi-format document content extraction."""

import csv
import io
import os
import re
from html.parser import HTMLParser
from pathlib import Path

from docwizard.config import SUPPORTED_READ_FORMATS, MAX_CONTENT_CHARS, MAX_FILE_SIZE_MB


class UnsupportedFormatError(Exception):
    pass


class FileTooLargeError(Exception):
    pass


def read_document(path: str) -> str:
    """Read document content as plain text.

    Raises UnsupportedFormatError for unknown formats,
    FileTooLargeError for files exceeding MAX_FILE_SIZE_MB.
    """
    p = Path(path)

    if not p.exists():
        raise FileNotFoundError(f"File not found: {path}")

    ext = p.suffix.lower()
    if ext not in SUPPORTED_READ_FORMATS:
        raise UnsupportedFormatError(
            f"Format '{ext}' is not supported. "
            f"Supported: {', '.join(sorted(SUPPORTED_READ_FORMATS))}"
        )

    # Check file size
    size_mb = p.stat().st_size / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        raise FileTooLargeError(
            f"File is {size_mb:.1f}MB, exceeds {MAX_FILE_SIZE_MB}MB limit."
        )

    # Specific readers for binary/complex formats
    specific_readers = {
        ".pdf": _read_pdf,
        ".docx": _read_docx,
        ".xlsx": _read_xlsx,
        ".pptx": _read_pptx,
        ".csv": _read_csv,
        ".html": _read_html,
        ".htm": _read_html,
        ".xhtml": _read_html,
        ".rtf": _read_rtf,
        ".doc": _read_legacy_office,
        ".xls": _read_legacy_office,
        ".ppt": _read_legacy_office,
        ".odt": _read_odf,
        ".ods": _read_odf,
        ".odp": _read_odf,
    }

    reader = specific_readers.get(ext, _read_plain_text)
    content = reader(path)

    # Truncate if too long
    if len(content) > MAX_CONTENT_CHARS:
        content = (
            content[:MAX_CONTENT_CHARS]
            + f"\n\n[Content truncated at {MAX_CONTENT_CHARS:,} characters. "
            "Ask about specific sections for more detail.]"
        )

    return content


# ── Plain text (catch-all for .txt, .md, .json, .xml, .py, .log, etc.) ──

def _read_plain_text(path: str) -> str:
    """Read any plain-text file."""
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(path, "r", encoding="latin-1", errors="replace") as f:
            return f.read()


# ── CSV ──────────────────────────────────────────────────────────────

def _read_csv(path: str) -> str:
    """Read CSV with summary + content."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    # Parse for summary
    try:
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
        if rows:
            header = rows[0]
            num_rows = len(rows) - 1
            summary = (
                f"CSV Summary: {num_rows} rows, {len(header)} columns\n"
                f"Columns: {', '.join(header)}\n\n"
            )
            return summary + content
    except csv.Error:
        pass

    return content


# ── HTML / HTM ───────────────────────────────────────────────────────

class _HTMLTextExtractor(HTMLParser):
    """Strip HTML tags and extract readable text."""

    def __init__(self):
        super().__init__()
        self._parts: list[str] = []
        self._skip = False
        self._skip_tags = {"script", "style", "head", "noscript", "svg"}

    def handle_starttag(self, tag, attrs):
        if tag.lower() in self._skip_tags:
            self._skip = True
        if tag.lower() in ("br", "p", "div", "li", "tr", "h1", "h2", "h3",
                           "h4", "h5", "h6", "blockquote", "pre"):
            self._parts.append("\n")

    def handle_endtag(self, tag):
        if tag.lower() in self._skip_tags:
            self._skip = False
        if tag.lower() in ("p", "div", "li", "tr", "h1", "h2", "h3",
                           "h4", "h5", "h6", "blockquote", "pre", "table"):
            self._parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self._parts.append(data)


def _read_html(path: str) -> str:
    """Read HTML file and extract text content."""
    raw = _read_plain_text(path)

    try:
        extractor = _HTMLTextExtractor()
        extractor.feed(raw)
        text = "".join(extractor._parts)
        # Clean up excessive whitespace
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if text:
            return text
    except Exception:
        pass

    # Fallback: return raw HTML
    return raw


# ── RTF ──────────────────────────────────────────────────────────────

def _read_rtf(path: str) -> str:
    """Read RTF file and extract text content."""
    raw = _read_plain_text(path)

    # Try striprtf if available
    try:
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(raw)
    except ImportError:
        pass

    # Basic RTF text extraction fallback
    # Remove RTF control words and groups
    text = re.sub(r"\\[a-z]+\d*\s?", " ", raw)
    text = re.sub(r"[{}]", "", text)
    text = re.sub(r"\\.", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    if text:
        return text

    return raw


# ── PDF ──────────────────────────────────────────────────────────────

def _read_pdf(path: str) -> str:
    """Read PDF using pdfplumber."""
    import pdfplumber

    text_parts = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(f"--- Page {i + 1} ---\n{page_text}")

    if not text_parts:
        return "[PDF contains no extractable text — may be image-based.]"

    return "\n\n".join(text_parts)


# ── DOCX ─────────────────────────────────────────────────────────────

def _read_docx(path: str) -> str:
    """Read DOCX using python-docx."""
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    if not paragraphs:
        return "[Document appears empty.]"
    return "\n\n".join(paragraphs)


# ── XLSX ─────────────────────────────────────────────────────────────

def _read_xlsx(path: str) -> str:
    """Read XLSX using openpyxl."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    text_parts = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_text = []
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if row_count >= 500:
                rows_text.append("[... truncated at 500 rows ...]")
                break
            cells = [str(c) if c is not None else "" for c in row]
            rows_text.append("\t".join(cells))
            row_count += 1

        if rows_text:
            text_parts.append(
                f"--- Sheet: {sheet_name} ({row_count} rows) ---\n"
                + "\n".join(rows_text)
            )

    wb.close()

    if not text_parts:
        return "[Workbook appears empty.]"

    return "\n\n".join(text_parts)


# ── PPTX ─────────────────────────────────────────────────────────────

def _read_pptx(path: str) -> str:
    """Read PPTX using python-pptx."""
    from pptx import Presentation

    prs = Presentation(path)
    text_parts = []

    for i, slide in enumerate(prs.slides):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        slide_texts.append(para.text)

        if slide_texts:
            text_parts.append(
                f"--- Slide {i + 1} ---\n" + "\n".join(slide_texts)
            )

    if not text_parts:
        return "[Presentation contains no extractable text.]"

    return "\n\n".join(text_parts)


# ── Legacy Office formats (.doc, .xls, .ppt) ────────────────────────

def _read_legacy_office(path: str) -> str:
    """Attempt to read legacy Office formats using textutil (macOS) or fallback."""
    import sys
    import subprocess

    ext = Path(path).suffix.lower()

    # macOS textutil can convert many formats to plain text
    if sys.platform == "darwin":
        try:
            result = subprocess.run(
                ["textutil", "-convert", "txt", "-stdout", path],
                capture_output=True,
                text=True,
                timeout=15,
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout
        except (subprocess.TimeoutExpired, FileNotFoundError, OSError):
            pass

    return (
        f"[Legacy {ext} format detected. For best results, save as "
        f"{'.docx' if ext == '.doc' else '.xlsx' if ext == '.xls' else '.pptx'} "
        "and reopen.]"
    )


# ── ODF formats (.odt, .ods, .odp) ──────────────────────────────────

def _read_odf(path: str) -> str:
    """Read OpenDocument formats by extracting content.xml from the ZIP."""
    import zipfile

    ext = Path(path).suffix.lower()

    try:
        with zipfile.ZipFile(path, "r") as zf:
            if "content.xml" in zf.namelist():
                raw_xml = zf.read("content.xml").decode("utf-8", errors="replace")
                # Strip XML tags to get text
                text = re.sub(r"<[^>]+>", " ", raw_xml)
                text = re.sub(r"\s+", " ", text).strip()
                if text:
                    return text
    except (zipfile.BadZipFile, KeyError, OSError):
        pass

    return f"[Could not extract text from {ext} file.]"
